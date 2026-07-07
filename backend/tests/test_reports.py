import pytest
import pdfplumber
from docx import Document as DocxDocument

from app.models.analysis import Analysis, RiskFinding
from app.models.document import Document
from app.models.enums import FileType, RiskCategory, RiskSeverity
from app.services.report_generator_service import ReportGeneratorService

pytestmark = pytest.mark.asyncio


def _build_fixtures():
    document = Document(
        id=1,
        owner_id=1,
        original_filename="test_contract.txt",
        stored_path="/tmp/does-not-matter.txt",
        file_type=FileType.TXT,
        file_size_bytes=100,
        sha256_hash="abc123",
    )
    analysis = Analysis(
        id=1,
        document_id=1,
        contract_type="Test Agreement",
        executive_summary="This is a test executive summary for report generation.",
        key_obligations=["Party A must pay on time.", "Party B must deliver goods."],
        payment_terms="Net 30 days",
        renewal_clause="Automatically renews annually.",
        confidentiality_clause="Both parties shall keep information confidential.",
        termination_clause="Either party may terminate with 30 days notice.",
        recommended_actions=["Negotiate a liability cap.", "Clarify renewal terms."],
        compliance_score=72.5,
        compliance_grade="B",
    )
    findings = [
        RiskFinding(
            id=1,
            analysis_id=1,
            category=RiskCategory.MISSING_CLAUSE,
            severity=RiskSeverity.MEDIUM,
            confidence=0.8,
            title="Missing Indemnification Clause",
            explanation="No indemnification clause was found.",
            supporting_clause_text=None,
            suggested_action="Add an indemnification clause.",
        ),
        RiskFinding(
            id=2,
            analysis_id=1,
            category=RiskCategory.LEGAL_RED_FLAG,
            severity=RiskSeverity.CRITICAL,
            confidence=0.95,
            title="Unlimited Liability Exposure",
            explanation="The contract does not cap liability.",
            supporting_clause_text="Party A shall be liable for all damages without limit.",
            suggested_action="Add a liability cap.",
        ),
    ]
    return document, analysis, findings


async def test_generate_pdf_report_contains_expected_sections(tmp_path, monkeypatch):
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "reports_dir", tmp_path)

    service = ReportGeneratorService()
    document, analysis, findings = _build_fixtures()
    path = await service.generate_pdf(document, analysis, findings)

    assert path.exists()
    with pdfplumber.open(path) as pdf:
        full_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert "AI Contract Risk Assessment Report" in full_text
    assert "Executive Summary" in full_text
    assert "Clause Analysis" in full_text
    assert "AI Risk Assessment" in full_text
    assert "Recommended Actions" in full_text
    assert "Unlimited Liability Exposure" in full_text


async def test_generate_docx_report_contains_expected_sections(tmp_path, monkeypatch):
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "reports_dir", tmp_path)

    service = ReportGeneratorService()
    document, analysis, findings = _build_fixtures()
    path = await service.generate_docx(document, analysis, findings)

    assert path.exists()
    doc = DocxDocument(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "AI Contract Risk Assessment Report" in full_text
    assert "Executive Summary" in full_text
    assert "Missing Indemnification Clause" in full_text
