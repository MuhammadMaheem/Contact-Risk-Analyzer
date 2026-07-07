import asyncio
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import get_settings
from app.exceptions import ReportGenerationError
from app.models.document import Document
from app.models.analysis import Analysis, RiskFinding

settings = get_settings()

_SEVERITY_HEX_PDF = {
    "critical": "#B91C1C",
    "high": "#C2410C",
    "medium": "#A16207",
    "low": "#4D7C0F",
}
_ACCENT_HEX = "#B45309"


class ReportGeneratorService:
    """Generates the AI risk assessment report as PDF (reportlab) or DOCX (python-docx).
    Every report includes: executive summary, clause analysis, risk assessment,
    recommendations."""

    async def generate_pdf(self, document: Document, analysis: Analysis, findings: list[RiskFinding]) -> Path:
        return await asyncio.to_thread(self._generate_pdf_sync, document, analysis, findings)

    async def generate_docx(self, document: Document, analysis: Analysis, findings: list[RiskFinding]) -> Path:
        return await asyncio.to_thread(self._generate_docx_sync, document, analysis, findings)

    def _output_path(self, document: Document, analysis: Analysis, extension: str) -> Path:
        filename = f"report_doc{document.id}_analysis{analysis.id}.{extension}"
        return settings.reports_dir / filename

    def _generate_pdf_sync(self, document: Document, analysis: Analysis, findings: list[RiskFinding]) -> Path:
        try:
            path = self._output_path(document, analysis, "pdf")
            doc = SimpleDocTemplate(str(path), pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "TitleAmber", parent=styles["Title"], textColor=colors.HexColor("#1C1917")
            )
            heading_style = ParagraphStyle(
                "HeadingAmber", parent=styles["Heading2"], textColor=colors.HexColor(_ACCENT_HEX)
            )
            body_style = styles["BodyText"]

            elements = [
                Paragraph("AI Contract Risk Assessment Report", title_style),
                Paragraph(f"Document: {document.original_filename}", body_style),
                Paragraph(f"Contract Type: {analysis.contract_type or 'Not identified'}", body_style),
                Paragraph(
                    f"Compliance Score: {analysis.compliance_score} "
                    f"(Grade {analysis.compliance_grade})",
                    body_style,
                ),
                Spacer(1, 0.25 * inch),
                Paragraph("Executive Summary", heading_style),
                Paragraph(analysis.executive_summary or "No summary available.", body_style),
                Spacer(1, 0.2 * inch),
                Paragraph("Key Obligations", heading_style),
            ]
            for obligation in analysis.key_obligations or []:
                elements.append(Paragraph(f"&bull; {obligation}", body_style))

            elements.append(Spacer(1, 0.2 * inch))
            elements.append(Paragraph("Clause Analysis", heading_style))
            clause_rows = [
                ["Clause", "Content"],
                ["Payment Terms", analysis.payment_terms or "Not found"],
                ["Renewal Clause", analysis.renewal_clause or "Not found"],
                ["Confidentiality Clause", analysis.confidentiality_clause or "Not found"],
                ["Termination Clause", analysis.termination_clause or "Not found"],
            ]
            clause_table = Table(clause_rows, colWidths=[1.7 * inch, 4.3 * inch])
            clause_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#292524")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D6D3D1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ]
                )
            )
            elements.append(clause_table)

            elements.append(Spacer(1, 0.25 * inch))
            elements.append(Paragraph("AI Risk Assessment", heading_style))
            for finding in findings:
                color_hex = _SEVERITY_HEX_PDF.get(finding.severity.value, "#1C1917")
                elements.append(
                    Paragraph(
                        f'<font color="{color_hex}"><b>[{finding.severity.value.upper()}] '
                        f"{finding.title}</b></font> (confidence: {finding.confidence:.0%})",
                        body_style,
                    )
                )
                elements.append(Paragraph(finding.explanation, body_style))
                if finding.supporting_clause_text:
                    elements.append(
                        Paragraph(f'<i>"{finding.supporting_clause_text}"</i>', body_style)
                    )
                elements.append(Spacer(1, 0.1 * inch))

            elements.append(Paragraph("Recommended Actions", heading_style))
            for action in analysis.recommended_actions or []:
                elements.append(Paragraph(f"&bull; {action}", body_style))

            doc.build(elements)
            return path
        except Exception as exc:  # noqa: BLE001
            raise ReportGenerationError(f"PDF report generation failed: {exc}") from exc

    def _generate_docx_sync(self, document: Document, analysis: Analysis, findings: list[RiskFinding]) -> Path:
        try:
            path = self._output_path(document, analysis, "docx")
            doc = DocxDocument()

            title = doc.add_heading("AI Contract Risk Assessment Report", level=0)
            title.runs[0].font.color.rgb = RGBColor(0x1C, 0x19, 0x17)

            doc.add_paragraph(f"Document: {document.original_filename}")
            doc.add_paragraph(f"Contract Type: {analysis.contract_type or 'Not identified'}")
            doc.add_paragraph(
                f"Compliance Score: {analysis.compliance_score} (Grade {analysis.compliance_grade})"
            )

            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph(analysis.executive_summary or "No summary available.")

            doc.add_heading("Key Obligations", level=1)
            for obligation in analysis.key_obligations or []:
                doc.add_paragraph(obligation, style="List Bullet")

            doc.add_heading("Clause Analysis", level=1)
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            header_cells = table.rows[0].cells
            header_cells[0].text = "Clause"
            header_cells[1].text = "Content"
            clause_data = [
                ("Payment Terms", analysis.payment_terms),
                ("Renewal Clause", analysis.renewal_clause),
                ("Confidentiality Clause", analysis.confidentiality_clause),
                ("Termination Clause", analysis.termination_clause),
            ]
            for label, content in clause_data:
                row = table.add_row().cells
                row[0].text = label
                row[1].text = content or "Not found"

            doc.add_heading("AI Risk Assessment", level=1)
            for finding in findings:
                p = doc.add_paragraph()
                run = p.add_run(f"[{finding.severity.value.upper()}] {finding.title} ")
                run.bold = True
                p.add_run(f"(confidence: {finding.confidence:.0%})")
                doc.add_paragraph(finding.explanation)
                if finding.supporting_clause_text:
                    quote = doc.add_paragraph()
                    quote_run = quote.add_run(f'"{finding.supporting_clause_text}"')
                    quote_run.italic = True

            doc.add_heading("Recommended Actions", level=1)
            for action in analysis.recommended_actions or []:
                doc.add_paragraph(action, style="List Bullet")

            doc.save(str(path))
            return path
        except Exception as exc:  # noqa: BLE001
            raise ReportGenerationError(f"DOCX report generation failed: {exc}") from exc


report_generator_service = ReportGeneratorService()
