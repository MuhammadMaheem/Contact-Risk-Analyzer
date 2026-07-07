import asyncio
import logging
from dataclasses import dataclass
from pathlib import Path

import docx
import pdfplumber
from pdf2image import convert_from_path

from app.exceptions import DocumentParsingError
from app.models.enums import FileType

logger = logging.getLogger(__name__)

OCR_TRIGGER_MIN_CHARS = 50


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None
    used_ocr: bool


class DocumentParserService:
    """Extracts plain text from PDF/DOCX/TXT, falling back to Tesseract OCR when a
    PDF has no usable text layer (i.e. it's a scanned image)."""

    async def extract_text(self, stored_path: str, file_type: FileType) -> ParsedDocument:
        try:
            if file_type == FileType.PDF:
                return await asyncio.to_thread(self._parse_pdf, stored_path)
            if file_type == FileType.DOCX:
                return await asyncio.to_thread(self._parse_docx, stored_path)
            return await asyncio.to_thread(self._parse_txt, stored_path)
        except DocumentParsingError:
            raise
        except Exception as exc:  # noqa: BLE001 - deliberately broad, converted to a domain error
            raise DocumentParsingError(f"Failed to parse document: {exc}") from exc

    def _parse_pdf(self, path: str) -> ParsedDocument:
        text_parts: list[str] = []
        page_count = 0
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        text = "\n\n".join(text_parts).strip()

        if len(text) >= OCR_TRIGGER_MIN_CHARS:
            return ParsedDocument(text=text, page_count=page_count, used_ocr=False)

        logger.info("PDF text layer insufficient (%d chars) — falling back to OCR", len(text))
        ocr_text = self._ocr_pdf(path)
        if len(ocr_text.strip()) < OCR_TRIGGER_MIN_CHARS:
            raise DocumentParsingError(
                "Could not extract meaningful text from this PDF, even with OCR fallback."
            )
        return ParsedDocument(text=ocr_text, page_count=page_count, used_ocr=True)

    def _ocr_pdf(self, path: str) -> str:
        import pytesseract

        images = convert_from_path(path)
        pages_text = [pytesseract.image_to_string(image) for image in images]
        return "\n\n".join(pages_text)

    def _parse_docx(self, path: str) -> ParsedDocument:
        document = docx.Document(path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if cell.text.strip()
        ]
        text = "\n\n".join(paragraphs + table_cells).strip()
        if len(text) < OCR_TRIGGER_MIN_CHARS:
            raise DocumentParsingError("DOCX file contains no extractable text.")
        return ParsedDocument(text=text, page_count=None, used_ocr=False)

    def _parse_txt(self, path: str) -> ParsedDocument:
        raw = Path(path).read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        text = text.strip()
        if len(text) < OCR_TRIGGER_MIN_CHARS:
            raise DocumentParsingError("Text file is empty or too short to analyze.")
        return ParsedDocument(text=text, page_count=None, used_ocr=False)


document_parser_service = DocumentParserService()
